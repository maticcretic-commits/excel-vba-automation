"""Tests for the demo workbook builder and the payroll math helpers.

The gross-pay / tax logic mirrors TimesheetPayroll.bas so the Python
side stays in lock-step with what the VBA produces.
"""

import datetime
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "python"))
import build_demo_workbook as builder

from openpyxl import load_workbook

# --- Mirror of the VBA payroll math ----------------------------------

DAILY_REGULAR_HOURS = 8
OVERTIME_MULTIPLIER = 1.5


def regular_hours(total):
    return min(total, DAILY_REGULAR_HOURS)


def overtime_hours(total):
    return max(0.0, total - DAILY_REGULAR_HOURS)


def gross_pay(regular, ot, rate):
    return regular * rate + ot * rate * OVERTIME_MULTIPLIER


def annual_tax(income):
    slabs = [
        (400000, 800000, 0.05),
        (800000, 1200000, 0.10),
        (1200000, 1600000, 0.15),
        (1600000, 2000000, 0.20),
        (2000000, 2400000, 0.25),
        (2400000, float("inf"), 0.30),
    ]
    tax = 0.0
    for lower, upper, rate in slabs:
        taxable = max(0.0, min(income, upper) - lower)
        tax += taxable * rate
    return tax


def monthly_tax(monthly_gross):
    return annual_tax(monthly_gross * 12) / 12


# --- Fixture ---------------------------------------------------------

@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    builder.OUT = tmp_path_factory.getbasetemp() / "demo_timesheet.xlsx"
    builder.main()
    return load_workbook(builder.OUT)


# --- Workbook structure tests ----------------------------------------

def test_all_expected_sheets_exist(workbook):
    for name in ("Employees", "Timesheet", "Data", "Lists",
                 "Dashboard", "Payroll", "Form"):
        assert name in workbook.sheetnames, f"missing sheet: {name}"


def test_employees_headers_and_rows(workbook):
    ws = workbook["Employees"]
    headers = [c.value for c in ws[1]]
    assert headers == ["Emp ID", "Name", "Department", "Role", "Hourly Rate"]
    ids = [r[0].value for r in ws.iter_rows(min_row=2)]
    assert ids == [f"E00{i}" for i in range(1, 9)]
    rates = [r[4].value for r in ws.iter_rows(min_row=2)]
    assert all(isinstance(x, (int, float)) and x > 0 for x in rates)


def test_timesheet_headers_and_sample_rows(workbook):
    ws = workbook["Timesheet"]
    headers = [c.value for c in ws[1]]
    assert headers == ["Emp ID", "Name", "Date", "Day", "Hours Worked"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) > 0
    for emp_id, name, dt, day, hours in rows:
        assert isinstance(dt, (datetime.date, datetime.datetime))
        assert dt.weekday() < 5, "timesheet should contain weekdays only"
        assert 0 <= hours <= 24


def test_data_sheet_structure(workbook):
    ws = workbook["Data"]
    headers = [c.value for c in ws[1]]
    assert headers == ["Date", "Category", "Metric", "Value"]
    values = [r[3].value for r in ws.iter_rows(min_row=2)]
    assert len(values) >= 40
    assert all(isinstance(v, (int, float)) for v in values)


def test_lists_sheet_headers_match_form_masters(workbook):
    ws = workbook["Lists"]
    headers = [c.value for c in ws[1]]
    assert headers == ["Department", "Role", "Shift"]
    for col in ws.iter_cols(min_row=2, max_row=ws.max_row, values_only=True):
        assert any(v is not None for v in col), "every list must have values"


# --- Payroll math edge cases -----------------------------------------

def test_zero_hours_means_zero_pay():
    assert gross_pay(regular_hours(0), overtime_hours(0), 500) == 0
    assert monthly_tax(0) == 0


def test_overtime_threshold_exactly_8_hours():
    assert regular_hours(8) == 8
    assert overtime_hours(8) == 0
    assert regular_hours(7.99) == pytest.approx(7.99)
    assert overtime_hours(8.01) == pytest.approx(0.01)


def test_overtime_premium_calculation():
    # 10-hour day at 500/hr: 8*500 + 2*500*1.5 = 5500
    assert gross_pay(regular_hours(10), overtime_hours(10), 500) == pytest.approx(5500)


def test_tax_slab_boundaries():
    assert annual_tax(400000) == 0                       # below first slab
    assert annual_tax(400001) == pytest.approx(0.05)     # first rupee taxed
    assert annual_tax(800000) == pytest.approx(20000)    # 5% of 4,00,000
    assert annual_tax(1200000) == pytest.approx(60000)   # +10% of 4,00,000
    assert annual_tax(3000000) == pytest.approx(480000)  # top slab reached


def test_monthly_tax_annualises():
    # monthly gross 1,00,000 -> annual 12,00,000 -> tax 60,000 -> 5,000/month
    assert monthly_tax(100000) == pytest.approx(5000)
    assert monthly_tax(30000) == 0  # annual 3,60,000 under the 4,00,000 threshold


def test_payroll_aggregation_demo_data(workbook):
    """Simulate CalculatePayroll over the demo timesheet for one employee."""
    ts = workbook["Timesheet"]
    emp = workbook["Employees"]
    rate = next(r[4] for r in emp.iter_rows(min_row=2, values_only=True)
                if r[0] == "E001")
    reg = ot = 0.0
    for row in ts.iter_rows(min_row=2, values_only=True):
        if row[0] == "E001":
            reg += regular_hours(row[4])
            ot += overtime_hours(row[4])
    gross = gross_pay(reg, ot, rate)
    tax = monthly_tax(gross)
    net = gross - tax
    assert reg > 0 and ot >= 0
    assert gross > tax >= 0
    assert net == pytest.approx(gross - tax)
