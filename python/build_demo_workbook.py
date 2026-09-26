"""Generate demo_timesheet.xlsx — a ready-made workbook the VBA macros can run on.

Sheets:
    Employees : A=Emp ID | B=Name | C=Department | D=Role | E=Hourly Rate
    Timesheet : A=Emp ID | B=Name | C=Date | D=Day | E=Hours Worked
    Data      : A=Date   | B=Category | C=Metric | D=Value
    Lists     : row 1 = master headers; values beneath each (for dropdowns)
    Dashboard : placeholder, rebuilt by OperationsDashboard.RefreshDashboard
    Payroll   : placeholder, rebuilt by TimesheetPayroll.CalculatePayroll
    Form      : B2 master choice / B3 dependent choice (labels only)
"""

import datetime
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = Path(__file__).resolve().parent.parent / "demo_timesheet.xlsx"

EMPLOYEES = [
    ("E001", "Aarav Sharma",   "Operations", "Analyst",    450),
    ("E002", "Diya Patel",     "Operations", "Supervisor", 650),
    ("E003", "Kabir Singh",    "Finance",    "Accountant", 550),
    ("E004", "Meera Iyer",     "Finance",    "Manager",    900),
    ("E005", "Rohan Verma",    "HR",         "Recruiter",  500),
    ("E006", "Ananya Gupta",   "HR",         "Manager",    850),
    ("E007", "Vikram Rao",     "IT",         "Developer",  750),
    ("E008", "Sneha Nair",     "IT",         "Lead",       1050),
]

CATEGORIES = ["Sales", "Support", "Logistics", "Marketing"]
METRICS = {
    "Sales":     ["Orders", "Revenue"],
    "Support":   ["Tickets Closed", "CSAT Responses"],
    "Logistics": ["Shipments", "On-time %"],
    "Marketing": ["Leads", "Campaign Clicks"],
}

LISTS = {
    "Department": ["Operations", "Finance", "HR", "IT"],
    "Role":       ["Analyst", "Supervisor", "Accountant", "Manager",
                   "Recruiter", "Developer", "Lead"],
    "Shift":      ["Morning", "Afternoon", "Night"],
}

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
THIN = Side(style="thin", color="C8C8C8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, ncols):
    for cell in ws[1]:
        if cell.column <= ncols:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = BORDER
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column <= ncols:
                cell.border = BORDER


def build_employees(wb):
    ws = wb.active
    ws.title = "Employees"
    ws.append(["Emp ID", "Name", "Department", "Role", "Hourly Rate"])
    for emp in EMPLOYEES:
        ws.append(list(emp))
    style_header(ws, 5)
    for col, width in zip("ABCDE", (10, 16, 14, 14, 14)):
        ws.column_dimensions[col].width = width


def build_timesheet(wb):
    ws = wb.create_sheet("Timesheet")
    ws.append(["Emp ID", "Name", "Date", "Day", "Hours Worked"])
    rng = random.Random(42)
    start = datetime.date(2026, 9, 1)
    for day in range(1, 23):  # first 22 days of Sep 2026
        d = start + datetime.timedelta(days=day - 1)
        if d.weekday() >= 5:
            continue
        for emp_id, name, *_ in EMPLOYEES:
            hours = rng.choice([7.5, 8, 8, 8, 8.5, 9, 10])
            ws.append([emp_id, name, d, d.strftime("%a"), hours])
    style_header(ws, 5)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = "dd-mmm-yyyy"
    for col, width in zip("ABCDE", (10, 16, 14, 8, 14)):
        ws.column_dimensions[col].width = width


def build_data(wb):
    ws = wb.create_sheet("Data")
    ws.append(["Date", "Category", "Metric", "Value"])
    rng = random.Random(7)
    start = datetime.date(2026, 8, 1)
    for i in range(60):
        d = start + datetime.timedelta(days=i)
        cat = rng.choice(CATEGORIES)
        metric = rng.choice(METRICS[cat])
        ws.append([d, cat, metric, rng.randint(20, 500)])
    style_header(ws, 4)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = "dd-mmm-yyyy"
    for col, width in zip("ABCD", (14, 14, 18, 10)):
        ws.column_dimensions[col].width = width


def build_lists(wb):
    ws = wb.create_sheet("Lists")
    headers = list(LISTS.keys())
    ws.append(headers)
    for i in range(max(len(v) for v in LISTS.values())):
        ws.append([LISTS[h][i] if i < len(LISTS[h]) else None for h in headers])
    style_header(ws, len(headers))
    for col in "ABC":
        ws.column_dimensions[col].width = 16


def build_placeholders(wb):
    dash = wb.create_sheet("Dashboard")
    dash["B2"] = "Run OperationsDashboard.RefreshDashboard to build this sheet."
    pay = wb.create_sheet("Payroll")
    pay["B2"] = "Run TimesheetPayroll.CalculatePayroll to build this sheet."
    form = wb.create_sheet("Form")
    form["A1"] = "Dropdown Demo"
    form["A2"] = "Category:"
    form["A3"] = "Value:"
    form["B2"] = "Department"
    for cell in ("A1", "A2", "A3"):
        form[cell].font = Font(bold=True)
    form["C2"] = "(B3 depends on B2 after running DropdownAutomation.BuildCascadingDropdowns)"
    form["C2"].font = Font(italic=True, color="787878")
    for ws in (dash, pay, form):
        for col in "ABC":
            ws.column_dimensions[col].width = 18


def main():
    wb = Workbook()
    build_employees(wb)
    build_timesheet(wb)
    build_data(wb)
    build_lists(wb)
    build_placeholders(wb)
    wb.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
